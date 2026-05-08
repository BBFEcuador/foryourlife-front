from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[1]
SRC_VIEWS = ROOT / "src" / "views"
OUTPUT_DIR = ROOT / "docs"
OUTPUT_FILE = OUTPUT_DIR / "Manual_Usuario_Foryourlife_Front.docx"


ROUTES = [
    {"name": "login-user", "path": "/login", "component": "src/views/authentication/LoginPage.vue", "role": "Publico", "permissions": []},
    {"name": "Authentication", "path": "/register/:token", "component": "src/views/authentication/ParticipantsRegister.vue", "role": "Publico", "permissions": []},
    {"name": "termsAndConditions", "path": "/terms-and-conditions", "component": "src/views/authentication/TermsAndConditions.vue", "role": "Publico", "permissions": []},
    {"name": "UserDashboard", "path": "/dashboard", "component": "src/views/user/UserDashboard.vue", "role": "Usuario", "permissions": []},
    {"name": "userPayments", "path": "/payments", "component": "src/views/user/PaymentsList.vue", "role": "Usuario", "permissions": []},
    {"name": "Admin Login", "path": "/admin/auth/login", "component": "src/views/admin/AdminLogin.vue", "role": "Administrador", "permissions": []},
    {"name": "Admin Reset Password", "path": "/admin/auth/reset-password", "component": "src/views/admin/AdminResetPassword.vue", "role": "Administrador", "permissions": []},
    {"name": "home-admin", "path": "/admin", "component": "src/views/admin/AdminDashboard.vue", "role": "Administrador", "permissions": []},
    {"name": "settings", "path": "/admin/settings", "component": "src/views/admin/AdminSettings.vue", "role": "Administrador", "permissions": ["SEE_CONFIGURATIONS"]},
    {"name": "users-admin", "path": "/admin/users", "component": "src/views/admin/users/AdminUsers.vue", "role": "Administrador", "permissions": ["SEE_USERS"]},
    {"name": "users-admin-add", "path": "/admin/users/add", "component": "src/views/admin/users/AddAdminUser.vue", "role": "Administrador", "permissions": ["CREATE_USERS"]},
    {"name": "sedes-admin", "path": "/admin/sedes", "component": "src/views/admin/sedes/AdminSedes.vue", "role": "Administrador", "permissions": ["SEE_CAMPUS"]},
    {"name": "training-admin", "path": "/admin/trainers", "component": "src/views/admin/training/AdminTraining.vue", "role": "Administrador", "permissions": ["SEE_TRAINERS"]},
    {"name": "programs-admin", "path": "/admin/program", "component": "src/views/admin/programs/AdminPrograms.vue", "role": "Administrador", "permissions": ["SEE_TRAININGS"]},
    {"name": "teams-admin", "path": "/admin/teams", "component": "src/views/admin/team/AdminTeams.vue", "role": "Administrador", "permissions": ["SEE_TEAMS"]},
    {"name": "teams-admin-add", "path": "/admin/teams/add", "component": "src/views/admin/team/AdminCreateTeamv2.vue", "role": "Administrador", "permissions": ["CREATE_TEAMS"]},
    {"name": "teams-admin-update", "path": "/admin/teams/edit/:id", "component": "src/views/admin/team/AdminTeamsUpdate.vue", "role": "Administrador", "permissions": ["UPDATE_TEAMS"]},
    {"name": "teams-admin-organigram", "path": "/admin/teams/organigram/:id", "component": "src/views/admin/team/AdminTeamsOrganigram.vue", "role": "Administrador", "permissions": ["UPDATE_TEAMS"]},
    {"name": "participants-admin", "path": "/admin/participants", "component": "src/views/admin/participants/AdminParticipants.vue", "role": "Administrador", "permissions": ["SEE_PARTICIPANTS"]},
    {"name": "participants-admin-edit", "path": "/admin/participants/edit/:id", "component": "src/views/admin/participants/EditParticipants.vue", "role": "Administrador", "permissions": ["UPDATE_PARTICIPANTS"]},
    {"name": "participants-admin-medical", "path": "/admin/participants/medical/:id", "component": "src/views/admin/participants/MedicalReportParticipants.vue", "role": "Administrador", "permissions": ["UPDATE_PARTICIPANTS"]},
    {"name": "participants-admin-contact-emergency", "path": "/admin/participants/contact-emergency/:id", "component": "src/views/admin/participants/ContactEmergencyParticipants.vue", "role": "Administrador", "permissions": ["UPDATE_PARTICIPANTS"]},
    {"name": "master-life-admin", "path": "/admin/master-life", "component": "src/views/admin/masterLife/AdminMasterLife.vue", "role": "Administrador", "permissions": ["SEE_MASTER_LIFES"]},
    {"name": "staff-admin", "path": "/admin/staff", "component": "src/views/admin/staff/AdminStaffList.vue", "role": "Administrador", "permissions": ["SEE_STAFF"]},
    {"name": "visionary-admin", "path": "/admin/visionarios", "component": "src/views/admin/visionaries/AdminVisionariesList.vue", "role": "Administrador", "permissions": ["SEE_VISIONARIES"]},
    {"name": "promote-focus", "path": "/admin/focusToYour/:id", "component": "src/views/admin/team/PromoteFocus.vue", "role": "Administrador", "permissions": []},
    {"name": "promote-your", "path": "/admin/yourToLife/:id", "component": "src/views/admin/team/PromoteYour.vue", "role": "Administrador", "permissions": []},
    {"name": "promote-life", "path": "/admin/lifeToMasterlife/:id", "component": "src/views/admin/team/PromoteLife.vue", "role": "Administrador", "permissions": []},
    {"name": "products-admin", "path": "/admin/products", "component": "src/views/admin/products/Products.vue", "role": "Administrador", "permissions": ["SEE_PRODUCTS"]},
    {"name": "discounts-admin", "path": "/admin/discounts", "component": "src/views/admin/discounts/DiscountList.vue", "role": "Administrador", "permissions": ["SEE_DISCOUNTS"]},
    {"name": "payments-admin", "path": "/admin/payments", "component": "src/views/admin/payments/PaymentsList.vue", "role": "Administrador", "permissions": ["SEE_PAYMENTS"]},
    {"name": "payments-admin-create", "path": "/admin/payments/create", "component": "src/views/admin/payments/PaymentsCreate.vue", "role": "Administrador", "permissions": ["SEE_PAYMENTS"]},
    {"name": "payment-methods", "path": "/admin/payment-methods", "component": "src/views/admin/paymentMethods/PaymentMethodsList.vue", "role": "Administrador", "permissions": ["SEE_PAYMENT_METHODS"]},
    {"name": "crm-admin", "path": "/admin/crm-admin", "component": "src/views/admin/crm/AdminCRM.vue", "role": "Administrador", "permissions": []},
    {"name": "reports", "path": "/admin/reports", "component": "src/views/admin/reports/AdminReports.vue", "role": "Administrador", "permissions": []},
    {"name": "operative-assistant-report", "path": "/admin/operative-assistant-report", "component": "src/views/admin/reports/OperationalAssistantReport.vue", "role": "Administrador", "permissions": []},
    {"name": "financial-administrator-report", "path": "/admin/financial-administrator-report", "component": "src/views/admin/reports/FinancialAdministratorReport.vue", "role": "Administrador", "permissions": []},
    {"name": "pos-main", "path": "/admin/pos", "component": "src/views/admin/pos/Pos.vue", "role": "Administrador", "permissions": ["SEE_EMISSION_POINTS"]},
    {"name": "cash-drawer-balances", "path": "/admin/pos/balances/:id", "component": "src/views/admin/pos/CashDrawersBalances.vue", "role": "Administrador", "permissions": []},
    {"name": "main-attendances", "path": "/admin/asistencias", "component": "src/views/admin/attendance/SelectAttendanceWithTraining.vue", "role": "Administrador", "permissions": ["SEE_ATTENDANCES_DECLARATIONS"]},
    {"name": "statements-admin", "path": "/admin/statements", "component": "src/views/admin/statements/StatementsList.vue", "role": "Administrador", "permissions": ["SEE_ATTENDANCES_DECLARATIONS"]},
    {"name": "trainer-login", "path": "/auth/trainer/login", "component": "src/views/trainer/login/TrainerLogin.vue", "role": "Entrenador", "permissions": []},
    {"name": "trainer-home", "path": "/trainer", "component": "src/views/trainer/dashboard/Dashboard.vue", "role": "Entrenador", "permissions": []},
    {"name": "trainer-team", "path": "/trainer/team/:id", "component": "src/views/trainer/team/TrainerTeam.vue", "role": "Entrenador", "permissions": []},
]


KEYWORD_PURPOSE = {
    "dashboard": "Muestra indicadores y accesos rapidos del modulo para la toma de decisiones operativas.",
    "login": "Gestiona la autenticacion del usuario y la apertura de sesion por perfil.",
    "reset": "Permite recuperar credenciales y restablecer contrasena de forma guiada.",
    "payment": "Administra informacion de cobros, pagos y trazabilidad financiera.",
    "discount": "Permite administrar reglas de descuento y su aplicacion en transacciones.",
    "product": "Gestiona catalogo de productos, configuraciones comerciales y disponibilidad.",
    "team": "Permite administrar equipos, composicion y cambios en su estructura.",
    "participant": "Permite consultar y mantener datos de participantes y su historial.",
    "trainer": "Centraliza la operacion relacionada con entrenadores y su seguimiento.",
    "program": "Administra entrenamientos, fechas, planeacion y agenda academica.",
    "attendance": "Gestiona asistencias, declaraciones y validaciones de cumplimiento.",
    "statement": "Consolida declaraciones y estados asociados al proceso de entrenamiento.",
    "report": "Presenta analitica e informes operativos para seguimiento y control.",
    "setting": "Permite parametrizar opciones globales del sistema.",
    "campus": "Gestiona sedes y configuracion de ubicaciones operativas.",
    "staff": "Administra personal de apoyo y su disponibilidad operativa.",
    "visionar": "Administra visionarios y su participacion en procesos internos.",
    "master": "Gestiona ciclo Master Life y su relacion con equipos y participantes.",
    "crm": "Permite gestionar relacion comercial y seguimiento de clientes.",
    "pos": "Permite operar punto de venta, cajas y procesos de recaudo.",
    "terms": "Presenta politicas, terminos y condiciones para aceptacion del usuario.",
    "register": "Gestiona el flujo de registro de participantes por invitacion o token.",
    "invitation": "Gestiona invitaciones y generacion de acceso por QR o enlace.",
    "error": "Pantalla de contingencia para rutas inexistentes o acceso no valido.",
}


def humanize_stem(stem: str) -> str:
    name = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", stem)
    name = name.replace("v2", " V2").replace("CRM", "CRM").replace("POS", "POS")
    return " ".join(name.split())


def infer_role(path: str) -> str:
    p = path.lower()
    if "/views/admin/" in p:
        return "Administrador"
    if "/views/trainer/" in p:
        return "Entrenador"
    if "/views/user/" in p:
        return "Usuario"
    if "/views/authentication/" in p:
        return "Publico"
    return "Soporte/Plantilla"


def infer_purpose(path: str) -> str:
    p = path.lower()
    for k, text in KEYWORD_PURPOSE.items():
        if k in p:
            return text
    return "Vista de soporte funcional utilizada para visualizacion, filtro o asistencia a un proceso del sistema."


def infer_actions(path: str) -> list[str]:
    p = path.lower()
    actions = ["Consultar informacion disponible en pantalla."]
    if any(x in p for x in ["list", "dashboard", "report", "view"]):
        actions.append("Filtrar y navegar registros segun criterios del usuario.")
    if any(x in p for x in ["create", "add", "register", "form"]):
        actions.append("Crear nuevos registros mediante formularios validados.")
    if any(x in p for x in ["edit", "update", "medical", "contact"]):
        actions.append("Editar informacion existente y guardar cambios controlados.")
    if any(x in p for x in ["login", "reset"]):
        actions.append("Gestionar autenticacion, recuperacion de acceso y cierre de sesion.")
    if any(x in p for x in ["pos", "payment", "cash"]):
        actions.append("Registrar operaciones economicas y validar estado de caja/pago.")
    if any(x in p for x in ["team", "promote", "organigram"]):
        actions.append("Gestionar estructura de equipos y promociones entre niveles.")
    if len(actions) == 1:
        actions.append("Ejecutar acciones de apoyo del flujo principal del modulo.")
    return actions


def extract_imports(view_path: Path) -> dict[str, list[str]]:
    content = view_path.read_text(encoding="utf-8", errors="ignore")
    imports = re.findall(r"import\s+.*?from\s+['\"](.*?)['\"]", content)
    data: dict[str, list[str]] = defaultdict(list)
    for imp in imports:
        if imp.startswith("@/components"):
            data["components"].append(imp)
        elif imp.startswith("@/composables"):
            data["composables"].append(imp)
        elif imp.startswith("@/stores"):
            data["stores"].append(imp)
        elif imp.startswith("@/api"):
            data["api"].append(imp)
    return data


def build_route_index() -> dict[str, list[dict[str, object]]]:
    index: dict[str, list[dict[str, object]]] = defaultdict(list)
    for route in ROUTES:
        index[str(route["component"])].append(route)
    return index


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def run() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()

    route_index = build_route_index()
    views = sorted(SRC_VIEWS.rglob("*.vue"))

    doc = Document()

    add_title(doc, "Manual de Usuario")
    add_subtitle(doc, "Proyecto: foryourlife-front")
    add_subtitle(doc, "Version del documento: 1.0")
    add_subtitle(doc, f"Fecha de emision: {today}")
    doc.add_page_break()

    doc.add_heading("1. Proposito del manual", level=1)
    doc.add_paragraph(
        "Este manual describe el uso funcional del sistema foryourlife-front, incluyendo vistas por perfil, flujos de operacion, permisos y recomendaciones de uso."
    )

    doc.add_heading("2. Alcance", level=1)
    doc.add_paragraph(
        "El documento cubre todas las vistas ubicadas en src/views y su relacion con rutas activas configuradas en el enrutador principal."
    )

    doc.add_heading("3. Perfiles de usuario", level=1)
    for role in ["Publico", "Usuario", "Entrenador", "Administrador", "Soporte/Plantilla"]:
        doc.add_paragraph(role, style="List Bullet")

    doc.add_heading("4. Arquitectura funcional resumida", level=1)
    doc.add_paragraph("Frontend: Vue 3 + TypeScript + Vite + Vuetify + Tailwind.")
    doc.add_paragraph("Estado: Pinia con persistencia local.")
    doc.add_paragraph("Seguridad: token Bearer, control por permisos y CASL ability.")
    doc.add_paragraph("Datos: Axios con interceptores por perfil (admin, trainer, user).")

    doc.add_heading("5. Matriz de permisos principales", level=1)
    for perm in [
        "SEE_USERS", "SEE_CAMPUS", "SEE_PRODUCTS", "SEE_DISCOUNTS", "SEE_PAYMENTS",
        "SEE_PAYMENT_METHODS", "SEE_TRAINERS", "SEE_TRAININGS", "SEE_TEAMS", "SEE_PARTICIPANTS",
        "SEE_MASTER_LIFES", "SEE_VISIONARIES", "SEE_STAFF", "SEE_ATTENDANCES_DECLARATIONS",
        "SEE_EMISSION_POINTS", "SEE_CONFIGURATIONS", "SEE_TRAINER_DASHBOARD", "SEE_OPERATIVE_DASHBOARD"
    ]:
        doc.add_paragraph(perm, style="List Bullet")

    doc.add_heading("6. Navegacion por rol", level=1)
    doc.add_paragraph("Administrador: Dashboard, Punto de Venta, Administrativo, Entrenamiento, Reportes y Configuracion.")
    doc.add_paragraph("Usuario: Invitacion y Mis pagos.")
    doc.add_paragraph("Entrenador: Dashboard y gestion de equipo asignado.")

    doc.add_heading("7. Catalogo detallado de vistas", level=1)
    doc.add_paragraph(
        "Cada ficha incluye objetivo, acceso, acciones esperadas, dependencias y riesgos operativos para soporte funcional."
    )

    for view in views:
        rel = view.relative_to(ROOT).as_posix()
        stem_name = humanize_stem(view.stem)
        role = infer_role("/" + rel)
        route_data = route_index.get(rel, [])
        imports = extract_imports(view)
        purpose = infer_purpose(rel)
        actions = infer_actions(rel)

        doc.add_heading(f"7.{views.index(view) + 1} {stem_name}", level=2)
        doc.add_paragraph(f"Archivo: {rel}")
        doc.add_paragraph(f"Perfil principal: {role}")
        doc.add_paragraph(f"Objetivo: {purpose}")

        if route_data:
            route_text = "; ".join(
                [f"{r['name']} ({r['path']})" for r in route_data]
            )
            doc.add_paragraph(f"Rutas asociadas: {route_text}")
            perm_set = sorted({p for r in route_data for p in r.get("permissions", [])})
            if perm_set:
                doc.add_paragraph("Permisos requeridos: " + ", ".join(perm_set))
            else:
                doc.add_paragraph("Permisos requeridos: No explicito en ruta (validacion por sesion y/o backend).")
        else:
            doc.add_paragraph("Rutas asociadas: No expuesta como ruta directa; se usa como subvista o plantilla de apoyo.")
            doc.add_paragraph("Permisos requeridos: Heredados de la vista contenedora.")

        doc.add_paragraph("Acciones del usuario:")
        for action in actions:
            doc.add_paragraph(action, style="List Bullet")

        doc.add_paragraph("Dependencias tecnicas relevantes:")
        if imports.get("components"):
            doc.add_paragraph("Componentes: " + ", ".join(sorted(set(imports["components"]))[:8]))
        if imports.get("composables"):
            doc.add_paragraph("Composables: " + ", ".join(sorted(set(imports["composables"]))[:8]))
        if imports.get("stores"):
            doc.add_paragraph("Stores: " + ", ".join(sorted(set(imports["stores"]))[:8]))
        if imports.get("api"):
            doc.add_paragraph("Servicios API: " + ", ".join(sorted(set(imports["api"]))[:8]))
        if not any(imports.values()):
            doc.add_paragraph("Dependencias locales: No se detectaron importaciones locales relevantes en el script de la vista.")

        doc.add_paragraph(
            "Consideraciones operativas: validar permisos del perfil, disponibilidad de token y consistencia de datos retornados por API antes de operar en produccion."
        )

    doc.add_heading("8. Flujos criticos recomendados para capacitacion", level=1)
    critical_flows = [
        "Ingreso al sistema por perfil (admin, entrenador, usuario).",
        "Operacion de punto de venta y control de caja.",
        "Gestion integral de equipos: creacion, edicion, organigrama y promociones.",
        "Gestion de participantes: edicion, ficha medica y contacto de emergencia.",
        "Ciclo de entrenamiento: entrenadores, programas, asistencias y declaraciones.",
        "Revision de reportes operativos y administrativos.",
    ]
    for flow in critical_flows:
        doc.add_paragraph(flow, style="List Number")

    doc.add_heading("9. Manejo de errores y soporte", level=1)
    doc.add_paragraph("Errores de autenticacion (401): cierre de sesion automatico y redireccion a login segun perfil.")
    doc.add_paragraph("Ruta invalida: visualizacion de pagina Error404Page.")
    doc.add_paragraph("Incidencias funcionales: revisar permisos del token y disponibilidad del servicio backend.")

    doc.add_heading("10. Buenas practicas de uso", level=1)
    practices = [
        "Operar siempre con el perfil correcto para evitar bloqueos por permisos.",
        "Confirmar sede/campus seleccionada antes de registrar operaciones.",
        "Evitar sesiones compartidas entre usuarios de distinto rol.",
        "Registrar datos completos en formularios para mantener trazabilidad.",
        "Consultar reportes para validacion de cierres operativos y financieros.",
    ]
    for practice in practices:
        doc.add_paragraph(practice, style="List Bullet")

    doc.add_heading("11. Glosario", level=1)
    glossary = [
        "Campus/Sede: unidad fisica de operacion.",
        "POS: punto de venta para cobros y control de caja.",
        "Declaraciones: registros de seguimiento de estado Focus/Your.",
        "Master Life: nivel avanzado del proceso de formacion.",
        "Permiso: autoridad funcional definida en el token del usuario.",
    ]
    for item in glossary:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUTPUT_FILE)
    print(f"Manual generado en: {OUTPUT_FILE}")
    print(f"Total de vistas documentadas: {len(views)}")


if __name__ == "__main__":
    run()
